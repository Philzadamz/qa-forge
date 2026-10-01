"""Extract plain text from an uploaded user story file (PRD §7.3, §8.2 step 1).

Image-based story pages (OCR via Claude vision, per the PRD) aren't implemented yet — that
needs a vision-capable provider call, which is out of scope while running on DeepSeek's
text-only chat API. Uploading an image raises `UnsupportedFileTypeError` rather than silently
producing empty text.
"""

import csv
import io
from pathlib import PurePosixPath

import docx
import openpyxl
import pdfplumber

TEXT_EXTENSIONS = {".txt", ".md"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".docx", ".pdf", ".xlsx", ".csv"}


class UnsupportedFileTypeError(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    suffix = PurePosixPath(filename).suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        return data.decode("utf-8", errors="replace")
    if suffix == ".docx":
        return _extract_docx(data)
    if suffix == ".pdf":
        return _extract_pdf(data)
    if suffix == ".xlsx":
        return _extract_xlsx(data)
    if suffix == ".csv":
        return _extract_csv(data)
    raise UnsupportedFileTypeError(
        f"Unsupported file type '{suffix}'. Supported: {sorted(SUPPORTED_EXTENSIONS)}"
    )


def _extract_docx(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pdf(data: bytes) -> str:
    parts = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                parts.append(text)
    return "\n\n".join(parts)


def _extract_xlsx(data: bytes) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    parts = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_csv(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    return "\n".join(
        " | ".join(cell.strip() for cell in row if cell.strip()) for row in reader if row
    )
