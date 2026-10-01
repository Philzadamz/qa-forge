"""Parse the DEFAULT SCENARIOS section out of a team-template test-case workbook (ADM-DC-3).

Reads rows between the `DEFAULT SCENARIOS` and `FUNCTIONAL SCENARIOS` banners, un-merges
the FEATURE column by propagating the feature name down merged ranges, and returns parsed
cases plus warnings (e.g. duplicate TESTCASE NO values) for the admin to review before saving.
Column positions are detected from the header row rather than hard-coded to Appendix B's
`A..H` letters, so a workbook with reordered columns still parses.
"""

import re
from dataclasses import dataclass, field
from io import BytesIO
from typing import cast

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

DEFAULT_BANNER = "DEFAULT SCENARIOS"
FUNCTIONAL_BANNER = "FUNCTIONAL SCENARIOS"
DEFAULT_EVIDENCE_GROUP = "Default Scenarios"

_STEP_NUMBER_RE = re.compile(r"^\s*\d+[.)]\s*")

_COLUMN_MATCHERS: dict[str, tuple[str, ...]] = {
    "id": ("TESTCASE",),
    "feature": ("FEATURE",),
    "scenario": ("OPERATION", "SCENARIO", "SCENERIO"),
    "steps": ("STEPS",),
    "expected": ("EXPECTED",),
    "actual": ("ACTUAL",),
    "status": ("STATUS",),
}


class XlsxImportError(ValueError):
    pass


@dataclass
class ParsedDefaultCase:
    feature: str
    scenario: str
    steps: list[str]
    expected_result: str
    default_actual_result: str
    evidence_group: str = DEFAULT_EVIDENCE_GROUP
    sort_order: int = 0


@dataclass
class ImportWarning:
    row: int | None
    message: str


@dataclass
class ImportResult:
    cases: list[ParsedDefaultCase] = field(default_factory=list)
    warnings: list[ImportWarning] = field(default_factory=list)


def _normalize(text: object) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(text or "").upper())


def _find_header_row(ws: Worksheet) -> int:
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 50)):
        for cell in row:
            if _normalize(cell.value).startswith("TESTCASE"):
                return cast(int, cell.row)
    raise XlsxImportError("Could not find the TESTCASE NO header row")


def _detect_columns(ws: Worksheet, header_row: int) -> dict[str, int]:
    columns: dict[str, int] = {}
    for cell in ws[header_row]:
        normalized = _normalize(cell.value)
        if not normalized:
            continue
        for key, needles in _COLUMN_MATCHERS.items():
            if key in columns:
                continue
            if any(needle in normalized for needle in needles):
                columns[key] = cast(int, cell.column)
    missing = {"feature", "scenario", "steps", "expected"} - columns.keys()
    if missing:
        raise XlsxImportError(f"Missing required column(s) in header row: {sorted(missing)}")
    return columns


def _find_banner_rows(ws: Worksheet, header_row: int) -> tuple[int, int]:
    default_row: int | None = None
    functional_row: int | None = None
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row, max_col=1):
        cell = row[0]
        normalized = _normalize(cell.value)
        if normalized == _normalize(DEFAULT_BANNER):
            default_row = cell.row
        elif normalized == _normalize(FUNCTIONAL_BANNER):
            functional_row = cell.row
            break
    if default_row is None:
        raise XlsxImportError(f"Could not find the '{DEFAULT_BANNER}' banner row")
    end_row = (functional_row - 1) if functional_row is not None else ws.max_row
    return default_row + 1, end_row


def _merged_column_lookup(ws: Worksheet, col: int) -> dict[int, object]:
    """Map every row in a vertically-merged range in `col` to the anchor cell's value.

    The template merges FEATURE across a feature's rows, and sometimes STEPS too when
    consecutive cases share identical steps (PRD §3.1) — both need un-merging the same way.
    """
    lookup: dict[int, object] = {}
    for merged_range in ws.merged_cells.ranges:
        if (
            merged_range.min_col <= col <= merged_range.max_col
            and merged_range.min_row != merged_range.max_row
        ):
            top_left = ws.cell(row=merged_range.min_row, column=col).value
            for r in range(merged_range.min_row, merged_range.max_row + 1):
                lookup[r] = top_left
    return lookup


def _split_steps(raw: object) -> list[str]:
    if not raw:
        return []
    lines = str(raw).replace("\r\n", "\n").split("\n")
    return [_STEP_NUMBER_RE.sub("", line).strip() for line in lines if line.strip()]


def parse_default_scenarios(file_bytes: bytes) -> ImportResult:
    try:
        wb = load_workbook(BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        raise XlsxImportError(f"Could not read workbook: {exc}") from exc

    ws = wb["Test Case"] if "Test Case" in wb.sheetnames else wb[wb.sheetnames[0]]
    header_row = _find_header_row(ws)
    columns = _detect_columns(ws, header_row)
    start_row, end_row = _find_banner_rows(ws, header_row)
    feature_lookup = _merged_column_lookup(ws, columns["feature"])
    steps_lookup = _merged_column_lookup(ws, columns["steps"])

    result = ImportResult()
    seen_ids: dict[str, int] = {}
    last_feature = ""
    sort_order = 0

    for r in range(start_row, end_row + 1):
        scenario = ws.cell(row=r, column=columns["scenario"]).value
        if scenario is None or not str(scenario).strip():
            continue

        feature = feature_lookup.get(r) or ws.cell(row=r, column=columns["feature"]).value
        feature = str(feature).strip() if feature else last_feature
        if not feature:
            result.warnings.append(ImportWarning(row=r, message="Row has no FEATURE value"))
        last_feature = feature or last_feature

        if "id" in columns:
            case_id = ws.cell(row=r, column=columns["id"]).value
            if case_id:
                case_id_str = str(case_id).strip()
                if case_id_str in seen_ids:
                    result.warnings.append(
                        ImportWarning(
                            row=r,
                            message=(
                                f"Duplicate TESTCASE NO '{case_id_str}' "
                                f"(first seen at row {seen_ids[case_id_str]})"
                            ),
                        )
                    )
                else:
                    seen_ids[case_id_str] = r

        expected = ws.cell(row=r, column=columns["expected"]).value
        actual = (
            ws.cell(row=r, column=columns.get("actual", 0)).value if "actual" in columns else None
        )
        steps_raw = steps_lookup.get(r) or ws.cell(row=r, column=columns["steps"]).value
        steps = _split_steps(steps_raw)
        if not steps:
            result.warnings.append(
                ImportWarning(row=r, message="Row has no STEPS TO EXECUTE value")
            )
            steps = ["(No steps found in the template — fill this in)"]

        expected_result = str(expected).strip() if expected else ""
        if not expected_result:
            result.warnings.append(ImportWarning(row=r, message="Row has no EXPECTED RESULT value"))
            expected_result = "(No expected result found in the template — fill this in)"

        result.cases.append(
            ParsedDefaultCase(
                feature=feature or "Uncategorised",
                scenario=str(scenario).strip(),
                steps=steps,
                expected_result=expected_result,
                default_actual_result=str(actual).strip() if actual else "As expected",
                sort_order=sort_order,
            )
        )
        sort_order += 1

    return result
