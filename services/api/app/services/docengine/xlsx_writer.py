"""Fill the team's QA_Test_Cases_Template.xlsx from suite data (PRD §9.1, Appendix B).

Rows 1-11 (header block + table column headers) are never touched beyond overwriting
value cells — they ship already styled exactly as the team wants, so re-styling them would
only risk drifting from the template. Rows 12+ (the old sample data) are cleared and
rewritten fresh, with styles captured from the template's own banner/body/id/evidence cells
before anything is deleted, so the output looks hand-filled rather than generated.
"""

import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image as PILImage

from app.core.enums import DEFAULT_STATUS_PALETTE, TestStatus
from app.services.docengine.models import ExportCase, SuiteExportData
from app.services.docengine.style import CellStyle

MAIN_SHEET = "Test Case"
DEFAULT_BANNER = "DEFAULT SCENARIOS"
FUNCTIONAL_BANNER = "FUNCTIONAL SCENARIOS"
TABLE_HEADER_ROW = 11
FIRST_DATA_ROW = 12
EVIDENCE_IMAGE_MAX_WIDTH_PX = 900

HEADER_CELLS = {
    "project_name_line": "B1",
    "user_group_dept": "B2",
    "solution_provider": "B3",
    "developers": "B4",
    "test_done_by": "B5",
    "test_reviewed_by": "B6",
    "start_date": "B7",
    "end_date": "B8",
    "test_description": "B9",
}
ENDPOINT_CELL = "E10"
RESULT_ANALYSIS_CELLS = {
    "cases_tested": "E2",
    "passes": "E3",
    "failures": "E4",
    "modification": "E5",
    "not_tested": "E6",
    "suspended": "E7",
    "regression": "E8",
}
REGRESSION_HELPER_COLUMN = "I"


def _argb(rgb_hex: str) -> str:
    """openpyxl needs 8-hex ARGB; a bare 6-hex string gets read back with a 00 (transparent)
    alpha rather than FF (opaque), so the status palette's plain RGB hexes need this."""
    return rgb_hex if len(rgb_hex) == 8 else f"FF{rgb_hex}"


def _sanitize_sheet_name(name: str, existing: set[str]) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", "", name).strip() or "Evidence"
    cleaned = cleaned[:31]
    candidate = cleaned
    suffix = 2
    while candidate in existing:
        trimmed = cleaned[: 31 - len(f" ({suffix})")]
        candidate = f"{trimmed} ({suffix})"
        suffix += 1
    return candidate


def _row_height_for(case: ExportCase) -> float:
    lines = max(len(case.steps), 1)
    return max(lines * 12.0, 15.0)


class _EvidenceSheets:
    """Lazily creates one evidence sheet per evidence group and tracks the write cursor."""

    def __init__(self, wb: Workbook) -> None:
        self._wb = wb
        self._sheet_names: dict[str, str] = {}
        self._next_row: dict[str, int] = {}

    def sheet_for(self, evidence_group: str) -> Worksheet:
        if evidence_group not in self._sheet_names:
            name = _sanitize_sheet_name(evidence_group, set(self._wb.sheetnames))
            self._sheet_names[evidence_group] = name
            self._next_row[name] = 1
            self._wb.create_sheet(name)
        return self._wb[self._sheet_names[evidence_group]]

    def sheet_name_for(self, evidence_group: str) -> str:
        self.sheet_for(evidence_group)  # ensure created
        return self._sheet_names[evidence_group]

    def place_evidence(self, case: ExportCase) -> str | None:
        """Write the case's evidence images and return the hyperlink display text, if any."""
        if not case.evidence:
            return None
        sheet = self.sheet_for(case.evidence_group)
        name = self._sheet_names[case.evidence_group]
        anchor_row = self._next_row[name]

        # PRD §3.1 specifies an en dash in the caption format, not a hyphen.
        sheet.cell(row=anchor_row, column=1, value=f"{case.case_no} – {case.scenario[:120]}")  # noqa: RUF001
        row_cursor = anchor_row + 1
        for image in case.evidence:
            with PILImage.open(image.file_path) as pil_img:
                width_px, height_px = pil_img.size
                if width_px > EVIDENCE_IMAGE_MAX_WIDTH_PX:
                    scale = EVIDENCE_IMAGE_MAX_WIDTH_PX / width_px
                    width_px, height_px = int(width_px * scale), int(height_px * scale)

            xl_img = XLImage(str(image.file_path))
            xl_img.width, xl_img.height = width_px, height_px
            sheet.add_image(xl_img, f"A{row_cursor}")
            row_cursor += max(1, height_px // 15) + 2

        self._next_row[name] = row_cursor + 1
        return f"'{name}'!A{anchor_row}"


def _write_case_row(
    ws: Worksheet,
    row: int,
    case: ExportCase,
    *,
    body_style: CellStyle,
    id_style: CellStyle,
    evidence_style: CellStyle,
    evidence_link: str | None,
) -> None:
    id_style.apply(ws.cell(row=row, column=1, value=case.case_no))
    body_style.apply(ws.cell(row=row, column=2, value=case.feature))
    body_style.apply(ws.cell(row=row, column=3, value=case.scenario))
    numbered_steps = "\n".join(f"{i}. {s}" for i, s in enumerate(case.steps, start=1))
    body_style.apply(ws.cell(row=row, column=4, value=numbered_steps))
    body_style.apply(ws.cell(row=row, column=5, value=case.expected_result))
    body_style.apply(ws.cell(row=row, column=6, value=case.actual_result))

    status_cell = ws.cell(row=row, column=7, value=case.status.value)
    body_style.apply(status_cell)
    palette = DEFAULT_STATUS_PALETTE.get(case.status)
    if palette:
        status_cell.fill = PatternFill("solid", fgColor=_argb(palette["fill"]))
        status_cell.font = Font(
            name=body_style.font.name,
            size=body_style.font.size,
            bold=True,
            color=_argb(palette["font"]),
        )

    evidence_cell = ws.cell(row=row, column=8)
    evidence_style.apply(evidence_cell)
    if evidence_link:
        sheet_ref = evidence_link.split("!", 1)[1]
        sheet_name = evidence_link.split("'!", 1)[0].lstrip("'")
        evidence_cell.value = evidence_link
        evidence_cell.hyperlink = Hyperlink(
            ref=evidence_cell.coordinate, location=f"'{sheet_name}'!{sheet_ref}"
        )

    helper_col = get_column_letter(9)
    ws[f"{helper_col}{row}"] = "Yes" if case.is_regression else "No"
    ws.column_dimensions[helper_col].hidden = True

    ws.row_dimensions[row].height = _row_height_for(case)


def _merge_feature_column(ws: Worksheet, start_row: int, end_row: int) -> None:
    if end_row < start_row:
        return
    run_start = start_row
    for row in range(start_row + 1, end_row + 2):
        same = (
            row <= end_row
            and ws.cell(row=row, column=2).value == ws.cell(row=run_start, column=2).value
        )
        if not same:
            if row - 1 > run_start:
                ws.merge_cells(start_row=run_start, start_column=2, end_row=row - 1, end_column=2)
            run_start = row


def write_test_cases_xlsx(template_path: Path, data: SuiteExportData, output_path: Path) -> None:
    wb = load_workbook(template_path)
    ws = wb[MAIN_SHEET]

    banner_style = CellStyle.capture(ws["A12"])
    body_style = CellStyle.capture(ws["C13"])
    id_style = CellStyle.capture(ws["A13"])
    evidence_style = CellStyle.capture(ws["H13"])

    # Drop the template's own sample evidence sheets — they belong to the sample suite.
    for name in list(wb.sheetnames):
        if name != MAIN_SHEET:
            del wb[name]

    # openpyxl's delete_rows doesn't reliably drop merged-cell ranges that lived entirely
    # within the deleted rows — they leak through as orphaned ranges. Unmerge explicitly first.
    for merged_range in list(ws.merged_cells.ranges):
        if merged_range.min_row >= FIRST_DATA_ROW:
            ws.unmerge_cells(str(merged_range))
    if ws.max_row >= FIRST_DATA_ROW:
        ws.delete_rows(FIRST_DATA_ROW, ws.max_row - FIRST_DATA_ROW + 1)

    for field, cell_ref in HEADER_CELLS.items():
        ws[cell_ref] = getattr(data.header, field)
    ws[ENDPOINT_CELL] = data.header.endpoint_url

    evidence_sheets = _EvidenceSheets(wb)
    row = FIRST_DATA_ROW

    def write_section(banner_text: str, cases: list[ExportCase]) -> tuple[int, int]:
        nonlocal row
        ws.cell(row=row, column=1, value=banner_text)
        banner_style.apply(ws.cell(row=row, column=1))
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        row += 1
        section_start = row
        for case in cases:
            link = evidence_sheets.place_evidence(case)
            _write_case_row(
                ws,
                row,
                case,
                body_style=body_style,
                id_style=id_style,
                evidence_style=evidence_style,
                evidence_link=link,
            )
            row += 1
        section_end = row - 1
        _merge_feature_column(ws, section_start, section_end)
        return section_start, section_end

    default_start, default_end = write_section(DEFAULT_BANNER, data.default_cases)
    functional_start, functional_end = write_section(FUNCTIONAL_BANNER, data.functional_cases)

    first_row = default_start if data.default_cases else functional_start
    last_row = functional_end if data.functional_cases else default_end
    if data.default_cases or data.functional_cases:
        g_range = f"G{first_row}:G{last_row}"
        i_range = f"I{first_row}:I{last_row}"
        ws[RESULT_ANALYSIS_CELLS["cases_tested"]] = f'=COUNTIF({g_range},"?*")'
        ws[RESULT_ANALYSIS_CELLS["passes"]] = f'=COUNTIF({g_range},"{TestStatus.PASSED.value}")'
        ws[RESULT_ANALYSIS_CELLS["failures"]] = f'=COUNTIF({g_range},"{TestStatus.FAILED.value}")'
        ws[RESULT_ANALYSIS_CELLS["modification"]] = (
            f'=COUNTIF({g_range},"{TestStatus.MODIFICATION.value}")'
        )
        ws[RESULT_ANALYSIS_CELLS["not_tested"]] = (
            f'=COUNTIF({g_range},"{TestStatus.NOT_TESTED.value}")'
        )
        ws[RESULT_ANALYSIS_CELLS["suspended"]] = (
            f'=COUNTIF({g_range},"{TestStatus.SUSPENDED.value}")'
        )
        ws[RESULT_ANALYSIS_CELLS["regression"]] = f'=COUNTIF({i_range},"Yes")'

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
