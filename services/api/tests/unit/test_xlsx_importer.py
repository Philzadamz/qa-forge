"""ADM-DC-3: parsing the DEFAULT SCENARIOS section out of a template-format workbook."""

from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.services.docengine.xlsx_importer import XlsxImportError, parse_default_scenarios

REAL_TEMPLATE = (
    Path(__file__).resolve().parents[4] / "templates" / "source" / "QA_Test_Cases_Template.xlsx"
)


def _build_workbook(rows: list[tuple[str | None, ...]], merges: list[str] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Case"
    headers = (
        "TESTCASE NO",
        "FEATURE",
        "OPERATIONS / SCENERIOS",
        "STEPS TO EXECUTE",
        "EXPECTED RESULT",
        "ACTUAL RESULT",
        "STATUS",
        "EVIDENCE",
    )
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for merge in merges or []:
        ws.merge_cells(merge)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_parses_default_section_between_banners() -> None:
    data = _build_workbook(
        [
            ("DEFAULT SCENARIOS", None, None, None, None, None, None, None),
            (
                "P_001",
                "Login",
                "Check that valid credentials grant access",
                "1. Open the app.\n2. Log in.",
                "User is logged in",
                "As expected",
                "Passed",
                None,
            ),
            (
                "P_002",
                "Login",
                "Check that invalid credentials are rejected",
                "1. Open the app.\n2. Enter bad credentials.",
                "Error shown",
                "As expected",
                "Passed",
                None,
            ),
            ("FUNCTIONAL SCENARIOS", None, None, None, None, None, None, None),
            (
                "P_003",
                "Some Feature",
                "This should not be imported",
                "1. Do a thing.",
                "Something happens",
                None,
                "Not Tested",
                None,
            ),
        ],
        # Row 1 = headers, row 2 = the DEFAULT SCENARIOS banner, rows 3-4 = the two Login
        # cases — merge their FEATURE cells the way the real template does.
        merges=["B3:B4"],
    )

    result = parse_default_scenarios(data)

    assert [c.scenario for c in result.cases] == [
        "Check that valid credentials grant access",
        "Check that invalid credentials are rejected",
    ]
    assert result.cases[0].feature == "Login"
    assert result.cases[1].feature == "Login"
    assert result.cases[0].steps == ["Open the app.", "Log in."]
    assert result.cases[0].evidence_group == "Default Scenarios"
    assert result.warnings == []


def test_merged_steps_column_is_propagated_like_feature() -> None:
    """Real template rows (e.g. D24:D26) merge STEPS across cases that share identical steps
    (PRD §3.1) — this used to crash the import with a 500 because only FEATURE was un-merged."""
    data = _build_workbook(
        [
            ("DEFAULT SCENARIOS", None, None, None, None, None, None, None),
            (
                "P_001",
                "Session Timeout",
                "Check that an idle session times out after the limit",
                "1. Log in.\n2. Wait for the timeout.\n3. Observe the session end.",
                "Session ends and user is redirected to login",
                None,
                None,
                None,
            ),
            (
                "P_002",
                "Session Timeout",
                "Check that activity resets the timeout",
                None,
                "Timer resets",
                None,
                None,
                None,
            ),
        ],
        # Row 1 = headers, row 2 = the DEFAULT SCENARIOS banner, rows 3-4 = the two cases —
        # merge their STEPS cells the way the real template does.
        merges=["D3:D4"],
    )

    result = parse_default_scenarios(data)

    assert result.cases[1].steps == [
        "Log in.",
        "Wait for the timeout.",
        "Observe the session end.",
    ]
    assert result.warnings == []


def test_missing_steps_warns_instead_of_crashing() -> None:
    data = _build_workbook(
        [
            ("DEFAULT SCENARIOS", None, None, None, None, None, None, None),
            ("P_001", "Login", "Check something", None, "Expected thing", None, None, None),
        ]
    )

    result = parse_default_scenarios(data)

    assert len(result.cases) == 1
    assert result.cases[0].steps  # never empty — DefaultTestCaseIn requires min_length=1
    assert any("STEPS" in w.message for w in result.warnings)


def test_missing_expected_result_warns_instead_of_crashing() -> None:
    data = _build_workbook(
        [
            ("DEFAULT SCENARIOS", None, None, None, None, None, None, None),
            ("P_001", "Login", "Check something", "1. Do it.", None, None, None, None),
        ]
    )

    result = parse_default_scenarios(data)

    assert len(result.cases) == 1
    assert result.cases[0].expected_result
    assert any("EXPECTED RESULT" in w.message for w in result.warnings)


def test_propagates_feature_without_merge_when_blank() -> None:
    data = _build_workbook(
        [
            ("DEFAULT SCENARIOS", None, None, None, None, None, None, None),
            (
                "P_001",
                "2FA",
                "Check OTP is required",
                "1. Log in.",
                "OTP prompt shown",
                None,
                None,
                None,
            ),
            ("P_002", None, "Check OTP expires", "1. Wait.", "OTP rejected", None, None, None),
        ]
    )

    result = parse_default_scenarios(data)

    assert [c.feature for c in result.cases] == ["2FA", "2FA"]


def test_warns_on_duplicate_testcase_no() -> None:
    data = _build_workbook(
        [
            ("DEFAULT SCENARIOS", None, None, None, None, None, None, None),
            ("P_001", "Login", "Check A", "1. Step.", "Result A", None, None, None),
            ("P_001", "Login", "Check B", "1. Step.", "Result B", None, None, None),
        ]
    )

    result = parse_default_scenarios(data)

    assert len(result.cases) == 2
    assert any("Duplicate" in w.message for w in result.warnings)


def test_missing_default_banner_raises() -> None:
    data = _build_workbook([("X_001", "Login", "Check A", "1. Step.", "Result", None, None, None)])

    with pytest.raises(XlsxImportError, match="DEFAULT SCENARIOS"):
        parse_default_scenarios(data)


def test_missing_required_column_raises() -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(("TESTCASE NO", "FEATURE"))
    ws.append(("DEFAULT SCENARIOS",))
    buffer = BytesIO()
    wb.save(buffer)

    with pytest.raises(XlsxImportError):
        parse_default_scenarios(buffer.getvalue())


@pytest.mark.skipif(
    not REAL_TEMPLATE.exists(), reason="real team template not present on this machine"
)
def test_parses_31_defaults_from_the_real_team_template() -> None:
    result = parse_default_scenarios(REAL_TEMPLATE.read_bytes())

    assert len(result.cases) == 31
    assert result.warnings == []
    feature_counts: dict[str, int] = {}
    for case in result.cases:
        feature_counts[case.feature] = feature_counts.get(case.feature, 0) + 1
    assert feature_counts["Login"] == 5
    assert feature_counts["2FA"] == 2
    assert feature_counts["Usability checks"] == 9
