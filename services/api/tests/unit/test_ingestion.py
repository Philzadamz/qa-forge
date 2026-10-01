import io

import docx
import openpyxl
import pytest

from app.services.ingestion.extract import UnsupportedFileTypeError, extract_text


def test_extract_txt() -> None:
    assert extract_text("story.txt", b"Check that login works.") == "Check that login works."


def test_extract_md() -> None:
    assert extract_text("story.md", b"# Story\n\nCheck login.") == "# Story\n\nCheck login."


def test_extract_docx() -> None:
    document = docx.Document()
    document.add_paragraph("Users must log in with email and password.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Limit"
    table.rows[0].cells[1].text = "150,000"
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_text("story.docx", buffer.getvalue())
    assert "Users must log in with email and password." in text
    assert "Limit | 150,000" in text


def test_extract_xlsx() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Feature", "Rule"])
    ws.append(["Login", "Max 5 attempts"])
    buffer = io.BytesIO()
    wb.save(buffer)

    text = extract_text("story.xlsx", buffer.getvalue())
    assert "Feature | Rule" in text
    assert "Login | Max 5 attempts" in text


def test_extract_csv() -> None:
    text = extract_text("story.csv", b"Feature,Rule\nLogin,Max 5 attempts\n")
    assert "Feature | Rule" in text
    assert "Login | Max 5 attempts" in text


def test_unsupported_extension_raises() -> None:
    with pytest.raises(UnsupportedFileTypeError):
        extract_text("story.png", b"\x89PNG...")
