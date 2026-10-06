from pathlib import Path

from docx import Document

from app.services.docengine.docx_writer import write_report_docx
from app.services.docengine.models import ApprovalRow, BugsSummary, ReportData, ResultAnalysis

TOKENIZED = (
    Path(__file__).resolve().parents[4] / "templates" / "tokenized" / "QA_Test_Report_Template.docx"
)


def _data(**overrides: object) -> ReportData:
    base = dict(
        product_name="Kusala",
        pr_links=[],
        version_numbers=[],
        test_url="",
        workitem_url="",
        jira_link="",
        general_description="",
        certified=False,
        features=[],
        exit_criteria=[],
        result_analysis=ResultAnalysis(
            test_cycles=1, total=0, passed=0, failed=0, unexecuted=0, suspended=0, modification=0
        ),
        automation_ratio="0%",
        bugs=BugsSummary(raised=0, fixed_retested=0, open=0),
        comments=["A comment."],
        approvals=[],
    )
    base.update(overrides)
    return ReportData(**base)  # type: ignore[arg-type]


def _render(tmp_path: Path, data: ReportData) -> Document:
    out = tmp_path / "report.docx"
    write_report_docx(TOKENIZED, data, out)
    return Document(out)


def test_empty_features_render_one_blank_row_with_unchecked_boxes(tmp_path: Path) -> None:
    doc = _render(tmp_path, _data())
    feature_body = doc.tables[2]
    assert len(feature_body.rows) == 1
    assert feature_body.rows[0].cells[3].text == "☐ Yes ☐ No"


def test_empty_approvals_render_one_blank_row(tmp_path: Path) -> None:
    doc = _render(tmp_path, _data())
    assert len(doc.tables[8].rows) == 2
    assert doc.tables[8].rows[1].cells[0].text == ""


def test_supplied_approvals_replace_the_defaults(tmp_path: Path) -> None:
    approvals = [ApprovalRow(action="Only", name="", staff_id="", signature="", date="")]
    doc = _render(tmp_path, _data(approvals=approvals))
    assert [row.cells[0].text for row in doc.tables[8].rows[1:]] == ["Only"]


def test_comments_render_inside_a_bordered_box_the_width_of_the_tables(tmp_path: Path) -> None:
    doc = _render(tmp_path, _data(comments=["One.", "Two."]))
    box = next(t for t in doc.tables if any("One." in c.text for r in t.rows for c in r.cells))
    assert len(box.rows) == 1 and len(box.columns) == 1
    assert "Two." in box.rows[0].cells[0].text
    assert "{{" not in box.rows[0].cells[0].text and "{%" not in box.rows[0].cells[0].text
    assert box._tbl.tblGrid.gridCol_lst[0].w == doc.tables[1]._tbl.tblGrid.gridCol_lst[
        0
    ].w * 0 + sum(g.w for g in doc.tables[1]._tbl.tblGrid.gridCol_lst)


def test_labels_are_tahoma_bold_and_values_are_verdana(tmp_path: Path) -> None:
    doc = _render(tmp_path, _data(product_name="Income Reversal"))
    project_table = doc.tables[0]
    label_run = project_table.rows[0].cells[0].paragraphs[0].runs[0]
    value_run = project_table.rows[0].cells[1].paragraphs[0].runs[0]
    assert label_run.font.name == "Tahoma" and label_run.font.bold is True
    assert value_run.font.name == "Verdana" and value_run.font.size.pt == 11
    assert value_run.text == "Income Reversal"
    header = doc.tables[1].rows[0].cells[1].paragraphs[0].runs[0]
    assert header.font.name == "Tahoma" and header.font.bold is True
