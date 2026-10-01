from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.golden.fixtures_paths import needs_xlsx_template


def _create_project(client: TestClient) -> dict:
    return client.post(
        "/api/v1/projects", json={"name": "Kusala", "app_code": "KUSALA", "id_prefix": "Kusala"}
    ).json()


def _create_type_with_one_default(admin_client: TestClient) -> dict:
    type_obj = admin_client.post("/api/v1/admin/types", json={"name": "Web Application"}).json()
    admin_client.post(
        f"/api/v1/admin/types/{type_obj['id']}/defaults",
        json={
            "feature": "Login",
            "scenario": "Check that valid credentials grant access",
            "steps": ["Open the app.", "Log in."],
            "expected_result": "User is logged in",
        },
    )
    return type_obj


@needs_xlsx_template
def test_export_xlsx_matches_template(admin_client: TestClient, tmp_path: Path) -> None:
    project = _create_project(admin_client)
    type_obj = _create_type_with_one_default(admin_client)
    suite = admin_client.post(
        "/api/v1/suites",
        json={
            "project_id": project["id"],
            "type_id": type_obj["id"],
            "name": "Export Suite",
            "header": {
                "project_name_line": "KUSALA: Export Suite",
                "user_group_dept": "Engineering",
                "test_done_by": "QA Engineer",
            },
        },
    ).json()

    resp = admin_client.post(f"/api/v1/suites/{suite['id']}/export/xlsx")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    out = tmp_path / "exported.xlsx"
    out.write_bytes(resp.content)
    wb = load_workbook(out)
    ws = wb["Test Case"]
    assert ws["B1"].value == "KUSALA: Export Suite"
    assert ws["A12"].value == "DEFAULT SCENARIOS"
    assert ws["A13"].value == "Kusala_001"
    assert ws["C13"].value == "Check that valid credentials grant access"
