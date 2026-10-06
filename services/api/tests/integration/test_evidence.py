import io
import uuid
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.enums import RunTarget
from app.models.evidence import Evidence
from app.models.test_run import TestRun
from app.services.storage import build_storage
from tests.golden.fixtures_paths import needs_xlsx_template


def _png_bytes() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (100, 50), color=(1, 2, 3)).save(buf, format="PNG")
    return buf.getvalue()


def _create_project(client: TestClient) -> dict:
    return client.post(
        "/api/v1/projects", json={"name": "Kusala", "app_code": "KUSALA", "id_prefix": "Kusala"}
    ).json()


def _suite_with_default_case(admin_client: TestClient) -> tuple[dict, dict]:
    project = _create_project(admin_client)
    type_obj = admin_client.post(
        "/api/v1/admin/types", json={"name": f"Evidence type {uuid.uuid4()}"}
    ).json()
    admin_client.post(
        f"/api/v1/admin/types/{type_obj['id']}/defaults",
        json={
            "feature": "Login",
            "scenario": "Check login",
            "steps": ["Log in."],
            "expected_result": "Logged in",
        },
    )
    suite = admin_client.post(
        "/api/v1/suites",
        json={
            "project_id": project["id"],
            "type_id": type_obj["id"],
            "name": "S",
            "header": {"project_name_line": "X"},
        },
    ).json()
    case = admin_client.get(f"/api/v1/suites/{suite['id']}/cases").json()[0]
    return suite, case


def _seed_test_lab_screenshot(db: Session, case: dict, suite: dict, caption: str) -> Evidence:
    run = TestRun(
        project_id=uuid.UUID(suite["project_id"]),
        suite_id=uuid.UUID(suite["id"]),
        target=RunTarget.WEB,
        selected_case_ids=[case["id"]],
    )
    db.add(run)
    db.flush()
    key = f"evidence/{suite['id']}/{case['id']}/{uuid.uuid4()}.png"
    build_storage(get_settings()).put(key, _png_bytes())
    row = Evidence(test_case_id=uuid.UUID(case["id"]), run_id=run.id, file_key=key, caption=caption)
    db.add(row)
    db.commit()
    return row


def test_manual_evidence_upload_is_removed(admin_client: TestClient) -> None:
    _, case = _suite_with_default_case(admin_client)
    resp = admin_client.post(
        f"/api/v1/cases/{case['id']}/evidence",
        files={"file": ("screenshot.png", _png_bytes(), "image/png")},
    )
    assert resp.status_code == 405


def test_listing_shows_only_test_lab_screenshots(
    admin_client: TestClient, db_session: Session
) -> None:
    suite, case = _suite_with_default_case(admin_client)
    _seed_test_lab_screenshot(db_session, case, suite, "Step 1")
    db_session.add(
        Evidence(
            test_case_id=uuid.UUID(case["id"]),
            run_id=None,
            file_key="evidence/legacy.png",
            caption="legacy manual upload",
        )
    )
    db_session.commit()

    listing = admin_client.get(f"/api/v1/cases/{case['id']}/evidence").json()
    assert [row["caption"] for row in listing] == ["Step 1"]


def test_screenshot_file_is_served(admin_client: TestClient, db_session: Session) -> None:
    suite, case = _suite_with_default_case(admin_client)
    row = _seed_test_lab_screenshot(db_session, case, suite, "Step 1")
    resp = admin_client.get(f"/api/v1/evidence/{row.id}/file")
    assert resp.status_code == 200
    assert resp.content == _png_bytes()


def test_evidence_cannot_be_deleted_by_hand(admin_client: TestClient, db_session: Session) -> None:
    suite, case = _suite_with_default_case(admin_client)
    row = _seed_test_lab_screenshot(db_session, case, suite, "Step 1")
    assert admin_client.delete(f"/api/v1/evidence/{row.id}").status_code == 404


@needs_xlsx_template
def test_export_puts_test_lab_screenshots_on_the_evidence_tabs(
    admin_client: TestClient, db_session: Session, tmp_path: Path
) -> None:
    suite, case = _suite_with_default_case(admin_client)
    _seed_test_lab_screenshot(db_session, case, suite, "Step 1")
    resp = admin_client.post(f"/api/v1/suites/{suite['id']}/export/xlsx")
    assert resp.status_code == 200, resp.text
    out = tmp_path / "export.xlsx"
    out.write_bytes(resp.content)

    wb = load_workbook(out)
    assert wb.sheetnames == ["Test Case", "Default test evidence", "Functional test evidence"]
    assert len(wb["Default test evidence"]._images) == 1
    assert len(wb["Functional test evidence"]._images) == 0
    link = wb["Test Case"]["H13"].hyperlink
    assert link is not None and "Default test evidence" in link.location


@needs_xlsx_template
def test_export_has_empty_evidence_tabs_when_nothing_was_run(
    admin_client: TestClient, tmp_path: Path
) -> None:
    suite, _ = _suite_with_default_case(admin_client)
    resp = admin_client.post(f"/api/v1/suites/{suite['id']}/export/xlsx")
    out = tmp_path / "export.xlsx"
    out.write_bytes(resp.content)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Test Case", "Default test evidence", "Functional test evidence"]
    assert all(len(wb[name]._images) == 0 for name in wb.sheetnames[1:])
