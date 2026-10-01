"""Golden-file tests for the docx writer and tokenizer (PRD §13.4)."""

from pathlib import Path

import docx
import pytest

from app.services.docengine.docx_tokenizer import DocxTokenizeError, tokenize_report_template
from app.services.docengine.docx_writer import write_report_docx
from app.services.docengine.models import (
    ApprovalRow,
    BugsSummary,
    ExceptionRow,
    FeatureRow,
    ReportData,
    ResultAnalysis,
)
from tests.golden.fixtures_paths import DOCX_TEMPLATE, needs_docx_template

TOKENIZED_TEMPLATE = (
    Path(__file__).resolve().parents[4] / "templates" / "tokenized" / "QA_Test_Report_Template.docx"
)
needs_tokenized_template = pytest.mark.skipif(
    not TOKENIZED_TEMPLATE.exists(), reason="tokenized template not built on this machine"
)


def _make_report_data(**overrides: object) -> ReportData:
    defaults: dict[str, object] = dict(
        product_name="Branch Petty Cash Retirement Flow on Kusala",
        pr_links=[
            "https://dev.azure.com/example/_git/repo1",
            "https://dev.azure.com/example/_git/repo2",
        ],
        version_numbers=["abc123", "def456"],
        test_url="http://example.com/dashboard",
        workitem_url="https://dev.azure.com/example/_workitems/edit/1",
        jira_link="https://example.atlassian.net/browse/ABC-1",
        general_description="Outlines the features of testing X sent to QA.",
        certified=True,
        features=[
            FeatureRow(name="Initiate Petty Cash Request", description="desc 1"),
            FeatureRow(
                name="Approve Request", description="desc 2", implemented=True, functional=False
            ),
        ],
        exit_criteria=["All cases executed", "All defects resolved", "UAT complete"],
        result_analysis=ResultAnalysis(
            test_cycles=1, total=60, passed=58, failed=2, unexecuted=0, suspended=0, modification=0
        ),
        automation_ratio="20:80",
        bugs=BugsSummary(raised=2, fixed_retested=2, open=0),
        exceptions=[
            ExceptionRow(
                case_id="Kusala_015",
                status_type="Failed",
                description="Something broke",
                severity="High",
                risk="Medium",
            )
        ],
        comments=["Functional testing has been completed.", "No blockers remain."],
        approvals=[
            ApprovalRow(action="Tested By", name="Jane Doe", staff_id="001", date="01/09/2026"),
            ApprovalRow(
                action="Product Owner Concurrence",
                name="John Smith",
                staff_id="002",
                date="02/09/2026",
            ),
        ],
        classification_label="Internal",
    )
    defaults.update(overrides)
    return ReportData(**defaults)  # type: ignore[arg-type]


@needs_docx_template
def test_tokenizer_runs_without_error(tmp_path: Path) -> None:
    out = tmp_path / "tokenized.docx"
    tokenize_report_template(DOCX_TEMPLATE, out)
    assert out.exists()


@needs_docx_template
def test_tokenizer_raises_clear_error_on_wrong_table_count(tmp_path: Path) -> None:
    import docx as docx_lib

    blank = tmp_path / "blank.docx"
    docx_lib.Document().save(blank)
    with pytest.raises(DocxTokenizeError, match="8 tables"):
        tokenize_report_template(blank, tmp_path / "out.docx")


@needs_tokenized_template
def test_render_leaves_no_jinja_tags(tmp_path: Path) -> None:
    out = tmp_path / "report.docx"
    write_report_docx(TOKENIZED_TEMPLATE, _make_report_data(), out)

    d = docx.Document(str(out))
    xml = d.element.xml
    assert "{{" not in xml
    assert "{%" not in xml


@needs_tokenized_template
def test_project_info_table_values(tmp_path: Path) -> None:
    out = tmp_path / "report.docx"
    write_report_docx(TOKENIZED_TEMPLATE, _make_report_data(), out)

    d = docx.Document(str(out))
    table = d.tables[0]
    assert table.rows[0].cells[1].text == "Branch Petty Cash Retirement Flow on Kusala"
    assert table.rows[1].cells[1].text == (
        "https://dev.azure.com/example/_git/repo1\nhttps://dev.azure.com/example/_git/repo2"
    )
    assert table.rows[2].cells[1].text == "abc123\ndef456"
    assert table.rows[6].cells[1].text == "Outlines the features of testing X sent to QA."


@needs_tokenized_template
def test_feature_table_row_count_and_checkboxes(tmp_path: Path) -> None:
    out = tmp_path / "report.docx"
    write_report_docx(TOKENIZED_TEMPLATE, _make_report_data(), out)

    d = docx.Document(str(out))
    table = d.tables[2]
    assert len(table.rows) == 2  # two features, no header (header lives in a separate table)
    assert table.rows[0].cells[0].text == "1."
    assert table.rows[0].cells[3].text == "☒ Yes ☐ No"  # implemented=True
    assert table.rows[1].cells[4].text == "☐ Yes ☒ No"  # functional=False


@needs_tokenized_template
def test_exit_criteria_row_count(tmp_path: Path) -> None:
    out = tmp_path / "report.docx"
    write_report_docx(TOKENIZED_TEMPLATE, _make_report_data(), out)

    d = docx.Document(str(out))
    table = d.tables[3]
    assert len(table.rows) == 4  # header + 3 criteria
    assert [r.cells[1].text for r in table.rows[1:]] == [
        "All cases executed",
        "All defects resolved",
        "UAT complete",
    ]


@needs_tokenized_template
def test_result_analysis_and_bugs_values(tmp_path: Path) -> None:
    out = tmp_path / "report.docx"
    write_report_docx(TOKENIZED_TEMPLATE, _make_report_data(), out)

    d = docx.Document(str(out))
    ra_table = d.tables[4]
    assert ra_table.rows[2].cells[1].text == "60"  # total
    assert ra_table.rows[3].cells[1].text == "58"  # passed

    ratio_table = d.tables[5]
    assert ratio_table.rows[0].cells[1].text == "20:80"
    assert ratio_table.rows[2].cells[1].text == "2"  # bugs raised


@needs_tokenized_template
def test_exceptions_row_for_actual_data(tmp_path: Path) -> None:
    out = tmp_path / "report.docx"
    write_report_docx(TOKENIZED_TEMPLATE, _make_report_data(), out)

    d = docx.Document(str(out))
    table = d.tables[6]
    assert len(table.rows) == 2  # header + 1 exception
    assert table.rows[1].cells[0].text == "Kusala_015"
    assert table.rows[1].cells[3].text == "High"


@needs_tokenized_template
def test_exceptions_na_row_when_empty(tmp_path: Path) -> None:
    out = tmp_path / "report.docx"
    write_report_docx(TOKENIZED_TEMPLATE, _make_report_data(exceptions=[]), out)

    d = docx.Document(str(out))
    table = d.tables[6]
    assert len(table.rows) == 2
    assert [c.text for c in table.rows[1].cells] == ["N/A", "N/A", "N/A", "N/A", "N/A"]


@needs_tokenized_template
def test_approvals_row_count(tmp_path: Path) -> None:
    out = tmp_path / "report.docx"
    write_report_docx(TOKENIZED_TEMPLATE, _make_report_data(), out)

    d = docx.Document(str(out))
    table = d.tables[7]
    assert len(table.rows) == 3  # header + 2 approvals
    assert table.rows[1].cells[1].text == "Jane Doe"
    assert table.rows[2].cells[1].text == "John Smith"


@needs_tokenized_template
def test_certificate_checkboxes_reflect_certified_flag(tmp_path: Path) -> None:
    out_yes = tmp_path / "yes.docx"
    write_report_docx(TOKENIZED_TEMPLATE, _make_report_data(certified=True), out_yes)
    out_no = tmp_path / "no.docx"
    write_report_docx(TOKENIZED_TEMPLATE, _make_report_data(certified=False), out_no)

    yes_doc = docx.Document(str(out_yes))
    no_doc = docx.Document(str(out_no))
    yes_text = next(p.text for p in yes_doc.paragraphs if "certified for deployment" in p.text)
    no_text = next(p.text for p in no_doc.paragraphs if "certified for deployment" in p.text)
    assert yes_text == "Is this build certified for deployment?\tYes ☒\tNo ☐"
    assert no_text == "Is this build certified for deployment?\tYes ☐\tNo ☒"


@needs_tokenized_template
def test_comments_render_as_separate_paragraphs(tmp_path: Path) -> None:
    out = tmp_path / "report.docx"
    write_report_docx(TOKENIZED_TEMPLATE, _make_report_data(), out)

    d = docx.Document(str(out))
    texts = [p.text for p in d.paragraphs if p.text.strip()]
    assert "Functional testing has been completed." in texts
    assert "No blockers remain." in texts


@needs_tokenized_template
def test_footer_classification_label(tmp_path: Path) -> None:
    out = tmp_path / "report.docx"
    write_report_docx(
        TOKENIZED_TEMPLATE, _make_report_data(classification_label="Confidential"), out
    )

    d = docx.Document(str(out))
    footer_texts = [
        t.text for t in d.sections[0].footer._element.iter() if t.tag.endswith("}t") and t.text
    ]
    assert footer_texts
    assert all(t == "Confidential" for t in footer_texts)
