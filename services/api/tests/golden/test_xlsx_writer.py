"""Golden-file tests for the xlsx writer (PRD §13.4)."""

import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PIL import Image as PILImage

from app.core.enums import TestStatus
from app.services.docengine.models import EvidenceImage, ExportCase, SuiteExportData, SuiteHeader
from app.services.docengine.xlsx_writer import write_test_cases_xlsx
from tests.golden.fixtures_paths import XLSX_TEMPLATE, needs_xlsx_template


def _make_header() -> SuiteHeader:
    return SuiteHeader(
        project_name_line="KUSALA: Branch Petty Cash Retirement Flow on Kusala",
        user_group_dept="Engineering / Technology",
        solution_provider="Sterling Bank / 1 Express",
        developers="Jane Doe, John Smith",
        test_done_by="QA Engineer",
        test_reviewed_by="QA Lead",
        start_date="01/09/2026",
        end_date="05/09/2026",
        test_description="This document seeks to outline the Branch Petty Cash Retirement Flow.",
        endpoint_url="http://kusala-fe-qa.example.com/dashboard",
    )


def _make_case(prefix: str, n: int, feature: str, status: TestStatus, **kw) -> ExportCase:
    return ExportCase(
        case_no=f"Kusala_{n:03d}",
        feature=feature,
        scenario=f"Check that {feature.lower()} scenario {n} behaves correctly",
        steps=["Do the first thing.", "Do the second thing."],
        expected_result="The expected thing happens",
        actual_result="As expected" if status == TestStatus.PASSED else "",
        status=status,
        **kw,
    )


@pytest.fixture
def evidence_png(tmp_path: Path) -> Path:
    path = tmp_path / "evidence.png"
    PILImage.new("RGB", (1200, 600), color=(10, 20, 30)).save(path)
    return path


@pytest.fixture
def sample_data(evidence_png: Path) -> SuiteExportData:
    defaults = [
        _make_case("Kusala", 1, "Login", TestStatus.PASSED),
        _make_case("Kusala", 2, "Login", TestStatus.PASSED),
        _make_case("Kusala", 3, "Logout", TestStatus.FAILED),
    ]
    functional = [
        _make_case(
            "Kusala",
            4,
            "Initiate Petty Cash Request",
            TestStatus.PASSED,
            evidence_group="Cash Hub",
            evidence=[EvidenceImage(file_path=evidence_png)],
        ),
        _make_case(
            "Kusala",
            5,
            "Initiate Petty Cash Request",
            TestStatus.NOT_TESTED,
            evidence_group="Cash Hub",
            is_regression=True,
        ),
        _make_case("Kusala", 6, "Approve Request", TestStatus.SUSPENDED, evidence_group="Cash Hub"),
    ]
    return SuiteExportData(
        header=_make_header(), default_cases=defaults, functional_cases=functional
    )


@needs_xlsx_template
def test_writes_header_block(sample_data: SuiteExportData, tmp_path: Path) -> None:
    out = tmp_path / "out.xlsx"
    write_test_cases_xlsx(XLSX_TEMPLATE, sample_data, out)

    wb = load_workbook(out)
    ws = wb["Test Case"]
    assert ws["B1"].value == sample_data.header.project_name_line
    assert ws["B4"].value == "Jane Doe, John Smith"
    assert ws["E10"].value == sample_data.header.endpoint_url
    # Header styling is untouched template styling, not re-applied by the writer.
    assert ws["A1"].fill.fgColor.rgb == "FFC00000"
    assert ws["B1"].fill.fgColor.rgb == "FFF4735E"


@needs_xlsx_template
def test_preserves_sheet_order_and_removes_sample_evidence_sheets(
    sample_data: SuiteExportData, tmp_path: Path
) -> None:
    out = tmp_path / "out.xlsx"
    write_test_cases_xlsx(XLSX_TEMPLATE, sample_data, out)

    wb = load_workbook(out)
    assert wb.sheetnames[0] == "Test Case"
    assert "Cash Center" not in wb.sheetnames  # sample evidence sheet from the template, gone
    assert "Cash Hub" in wb.sheetnames  # created fresh because a functional case needs it


@needs_xlsx_template
def test_banners_and_continuous_numbering(sample_data: SuiteExportData, tmp_path: Path) -> None:
    out = tmp_path / "out.xlsx"
    write_test_cases_xlsx(XLSX_TEMPLATE, sample_data, out)

    wb = load_workbook(out)
    ws = wb["Test Case"]
    assert ws["A12"].value == "DEFAULT SCENARIOS"
    assert "A12:H12" in [str(r) for r in ws.merged_cells.ranges]
    assert ws["A13"].value == "Kusala_001"
    assert ws["A14"].value == "Kusala_002"
    assert ws["A15"].value == "Kusala_003"
    assert ws["A16"].value == "FUNCTIONAL SCENARIOS"
    assert ws["A17"].value == "Kusala_004"
    assert ws["A19"].value == "Kusala_006"


@needs_xlsx_template
def test_feature_column_merges_within_section_not_across_banner(
    sample_data: SuiteExportData, tmp_path: Path
) -> None:
    out = tmp_path / "out.xlsx"
    write_test_cases_xlsx(XLSX_TEMPLATE, sample_data, out)

    wb = load_workbook(out)
    ws = wb["Test Case"]
    merges = [str(r) for r in ws.merged_cells.ranges]
    assert "B13:B14" in merges  # two consecutive "Login" defaults merged
    assert "B17:B18" in merges  # two consecutive "Initiate..." functional cases merged
    # Logout (row 15) is a single row, no merge; banner row 16 is never part of a B merge.
    assert not any(m.startswith("B15:") or "B16" in m for m in merges)


@needs_xlsx_template
def test_status_fills_match_palette(sample_data: SuiteExportData, tmp_path: Path) -> None:
    out = tmp_path / "out.xlsx"
    write_test_cases_xlsx(XLSX_TEMPLATE, sample_data, out)

    wb = load_workbook(out)
    ws = wb["Test Case"]
    assert ws["G13"].value == "Passed"
    assert ws["G13"].fill.fgColor.rgb == "FF00B050"
    assert ws["G15"].value == "Failed"
    assert ws["G15"].fill.fgColor.rgb == "FFFF0000"


@needs_xlsx_template
def test_result_analysis_formulas_are_correct_text(
    sample_data: SuiteExportData, tmp_path: Path
) -> None:
    out = tmp_path / "out.xlsx"
    write_test_cases_xlsx(XLSX_TEMPLATE, sample_data, out)

    wb = load_workbook(out)
    ws = wb["Test Case"]
    assert ws["E2"].value == '=COUNTIF(G13:G19,"?*")'
    assert ws["E3"].value == '=COUNTIF(G13:G19,"Passed")'
    assert ws["E4"].value == '=COUNTIF(G13:G19,"Failed")'
    assert ws["E8"].value == '=COUNTIF(I13:I19,"Yes")'


@needs_xlsx_template
def test_regression_helper_column_hidden_and_correct(
    sample_data: SuiteExportData, tmp_path: Path
) -> None:
    out = tmp_path / "out.xlsx"
    write_test_cases_xlsx(XLSX_TEMPLATE, sample_data, out)

    wb = load_workbook(out)
    ws = wb["Test Case"]
    assert ws.column_dimensions["I"].hidden is True
    assert ws["I18"].value == "Yes"  # Kusala_005 was marked is_regression
    assert ws["I13"].value == "No"


@needs_xlsx_template
def test_evidence_hyperlink_resolves_to_existing_cell_with_image(
    sample_data: SuiteExportData, tmp_path: Path
) -> None:
    out = tmp_path / "out.xlsx"
    write_test_cases_xlsx(XLSX_TEMPLATE, sample_data, out)

    wb = load_workbook(out)
    ws = wb["Test Case"]
    evidence_cell = ws["H17"]  # Kusala_004, the only case with an evidence image
    assert evidence_cell.value == "'Cash Hub'!A1"
    assert evidence_cell.hyperlink.location == "'Cash Hub'!A1"

    cash_hub = wb["Cash Hub"]
    assert cash_hub["A1"].value.startswith("Kusala_004")
    assert len(cash_hub._images) == 1


@needs_xlsx_template
def test_case_without_evidence_has_blank_evidence_cell(
    sample_data: SuiteExportData, tmp_path: Path
) -> None:
    out = tmp_path / "out.xlsx"
    write_test_cases_xlsx(XLSX_TEMPLATE, sample_data, out)

    wb = load_workbook(out)
    ws = wb["Test Case"]
    assert ws["H13"].value is None  # Kusala_001 has no evidence
    assert ws["H13"].hyperlink is None


@needs_xlsx_template
def test_font_and_column_widths_match_template(
    sample_data: SuiteExportData, tmp_path: Path
) -> None:
    out = tmp_path / "out.xlsx"
    write_test_cases_xlsx(XLSX_TEMPLATE, sample_data, out)

    wb = load_workbook(out)
    ws = wb["Test Case"]
    assert ws["C13"].font.name == "Century Gothic"
    assert ws["C13"].font.size == 9
    assert ws["C13"].alignment.wrap_text is True
    assert round(ws.column_dimensions["C"].width, 1) == 53.2


@needs_xlsx_template
def test_empty_suite_does_not_crash(tmp_path: Path) -> None:
    out = tmp_path / "out.xlsx"
    data = SuiteExportData(header=_make_header(), default_cases=[], functional_cases=[])
    write_test_cases_xlsx(XLSX_TEMPLATE, data, out)

    wb = load_workbook(out)
    ws = wb["Test Case"]
    assert ws["A12"].value == "DEFAULT SCENARIOS"


@pytest.mark.needs_libreoffice
@needs_xlsx_template
def test_formulas_recalculate_to_expected_counts(
    sample_data: SuiteExportData, tmp_path: Path
) -> None:
    if not shutil.which("soffice"):
        pytest.skip("LibreOffice not installed")
    import subprocess

    out = tmp_path / "out.xlsx"
    write_test_cases_xlsx(XLSX_TEMPLATE, sample_data, out)
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path), str(out)],
        check=True,
        timeout=60,
    )
    wb = load_workbook(out, data_only=True)
    ws = wb["Test Case"]
    assert ws["E2"].value == 6
    assert ws["E3"].value == 3
    assert ws["E4"].value == 1
